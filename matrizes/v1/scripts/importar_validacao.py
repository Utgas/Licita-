#!/usr/bin/env python3
"""Leva a revisão feita no xlsx (colunas validacao e comentario_validacao do M03 e do M04) de volta aos CSV.

Uso: python3 -I importar_validacao.py <matrizes_v1.xlsx devolvido> <pasta_Matrizes_v1> <saida_log.txt>

Lê as abas cujo nome começa por M03_ e M04_, casa as linhas pelo ID e grava só as duas colunas de revisão no CSV
(o resto da linha não muda). Valores fora de OK, CORRIGIR e DUVIDA são recusados e listados no log. Diferenças em outras
colunas entre o xlsx e o CSV são listadas como aviso e não são importadas. O log lista as linhas CORRIGIR e DUVIDA, que
viram trabalho na próxima rodada (a correção é feita no gerador da matriz, e as colunas de revisão são preservadas).
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comum import gravar_csv, ler_csv  # noqa: E402

DOMINIO = {"", "OK", "CORRIGIR", "DUVIDA"}


def main():
    xlsx, pasta, saida = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    log = []
    for prefixo in ("M03_", "M04_"):
        abas = [n for n in wb.sheetnames if n.startswith(prefixo)]
        if not abas:
            log.append(f"{prefixo}: aba não encontrada")
            continue
        ws = wb[abas[0]]
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(c) if c is not None else "" for c in linhas[0]]
        csv_path = next(pasta.glob(f"{prefixo}*.csv"))
        rows = ler_csv(csv_path)
        por_id = {r["id"]: r for r in rows}
        n, recusados, avisos, pendentes = 0, [], [], []
        for v in linhas[1:]:
            d = {cab[i]: ("" if x is None else str(x)) for i, x in enumerate(v) if i < len(cab)}
            rid = d.get("id", "")
            if rid not in por_id:
                avisos.append(f"ID {rid} não existe no CSV")
                continue
            val = d.get("validacao", "").strip().upper()
            if val not in DOMINIO:
                recusados.append(f"{rid}: '{val}'")
                continue
            for col in cab:
                if col not in ("validacao", "comentario_validacao") and col in por_id[rid] and \
                        d.get(col, "").replace("\r\n", "\n") != por_id[rid][col].replace("\r\n", "\n"):
                    avisos.append(f"{rid}.{col} editado no xlsx (não importado)")
            if val or d.get("comentario_validacao", ""):
                por_id[rid]["validacao"] = val
                por_id[rid]["comentario_validacao"] = d.get("comentario_validacao", "").strip()
                n += 1
                if val in ("CORRIGIR", "DUVIDA"):
                    pendentes.append(f"{rid} [{val}] {por_id[rid]['comentario_validacao']}")
        gravar_csv(csv_path, list(rows[0].keys()), rows, preservar=())
        log += [f"{csv_path.name}: {n} linhas com revisão importada; recusados {len(recusados)}; avisos {len(avisos)}"]
        log += [f"  RECUSADO {x}" for x in recusados] + [f"  AVISO {x}" for x in avisos[:200]]
        log += [f"  PENDENTE {x}" for x in pendentes]
    saida.write_text("\n".join(log) + "\n", encoding="utf-8")
    print("\n".join(log[:40]))


if __name__ == "__main__":
    main()
