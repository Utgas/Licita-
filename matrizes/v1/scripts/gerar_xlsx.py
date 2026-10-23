#!/usr/bin/env python3
"""Monta matrizes_v1.xlsx com uma aba por CSV (o CSV e a fonte).

Uso: python3 -I gerar_xlsx.py <pasta_Matrizes_v1>
Cada aba: cabecalho congelado, filtro automatico, texto com quebra de linha, larguras limitadas.
Colunas de revisao (validacao e comentario_validacao, no M03 e no M04): fundo amarelo e lista OK, CORRIGIR, DUVIDA.
Sao as unicas que o importar_validacao.py leva de volta aos CSV; edicao em outra coluna nao e importada.
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def main():
    pasta = Path(sys.argv[1])
    wb = Workbook()
    ins = wb.active
    ins.title = "LEIA_ME"
    for linha in ["Matrizes v1 do Licita+. O CSV é a fonte; esta planilha é cópia para leitura e revisão.",
                  "Revisão: preencha só as colunas amarelas 'validacao' (OK, CORRIGIR ou DUVIDA) e 'comentario_validacao' "
                  "nas abas M03 e M04.",
                  "Ao devolver a planilha, o script importar_validacao.py leva essas duas colunas para os CSV, pelo ID da linha.",
                  "Edições em outras colunas não são importadas: descreva a correção no comentario_validacao."]:
        ins.append([linha])
    ins.column_dimensions["A"].width = 130
    for arq in sorted(pasta.glob("M[0-9][0-9]_*.csv")):
        with open(arq, encoding="utf-8-sig", newline="") as f:
            linhas = list(csv.reader(f))
        ws = wb.create_sheet(title=arq.stem[:31])
        for r in linhas:
            ws.append(r)
        cab = Font(bold=True, color="FFFFFF")
        fundo = PatternFill("solid", fgColor="1F3A5F")
        for c in ws[1]:
            c.font = cab
            c.fill = fundo
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for col in range(1, len(linhas[0]) + 1):
            maior = max(len(r[col - 1]) if col - 1 < len(r) else 0 for r in linhas)
            ws.column_dimensions[get_column_letter(col)].width = max(10, min(60, maior * 0.9))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        cabecalho = linhas[0]
        if "comentario_validacao" in cabecalho:
            amarelo = PatternFill("solid", fgColor="FFF2CC")
            dv = DataValidation(type="list", formula1='"OK,CORRIGIR,DUVIDA"', allow_blank=True,
                                error="Use OK, CORRIGIR ou DUVIDA", errorTitle="validacao")
            ws.add_data_validation(dv)
            for nome in ("validacao", "comentario_validacao"):
                col = cabecalho.index(nome) + 1
                letra = get_column_letter(col)
                ws.column_dimensions[letra].width = 16 if nome == "validacao" else 45
                for row in range(2, len(linhas) + 1):
                    ws.cell(row=row, column=col).fill = amarelo
                if nome == "validacao":
                    dv.add(f"{letra}2:{letra}{len(linhas)}")
    destino = pasta / "matrizes_v1.xlsx"
    wb.save(destino)
    print(f"{destino}: {len(wb.sheetnames)} aba(s): {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
